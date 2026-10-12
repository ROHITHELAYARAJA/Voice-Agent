import asyncio
import io
import json
import time
import wave
import xml.etree.ElementTree as ET
from dataclasses import replace
from urllib.parse import urlsplit

import httpx
from openpyxl import load_workbook
from twilio.request_validator import RequestValidator

from buddy.logging import configure_logging
from buddy.runtime import BuddyRuntime
from config import ROOT, Settings
from main import create_app


async def verify(cloud: bool = False) -> None:
    output = ROOT / "outputs" / ("deepgram-verification" if cloud else "local-verification")
    output.mkdir(parents=True, exist_ok=True)
    settings = replace(
        Settings.from_env(),
        call_mode="gather",
        gather_ai_provider="deepgram" if cloud else "local",
        phone_voice_provider="eleven_labs" if cloud else "twilio",
        data_dir=output,
        excel_file=str(output / "Verified Check-in.xlsx"),
        public_hostname="https://verification.example.com",
    )
    configure_logging(settings.timezone)
    runtime = BuddyRuntime(settings)
    session = runtime.new_session()
    sid = "CA-local-verification-" + session["id"]
    runtime.store.update(session["id"], call_sid=sid, call_status="in-progress")
    durations = []
    audio_seconds = []
    played_audio = []
    try:
        await runtime.phone.check()
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app=create_app(settings, runtime)),
            base_url=settings.public_url,
        ) as client:

            async def send(path: str, speech: str = "") -> str:
                data = {
                    "AccountSid": settings.twilio_account_sid,
                    "CallSid": sid,
                    "To": settings.owner_phone_number,
                    "SpeechResult": speech,
                }
                signature = RequestValidator(settings.twilio_auth_token).compute_signature(
                    settings.public_url + path, data
                )
                started = time.monotonic()
                response = await client.post(
                    path, data=data, headers={"X-Twilio-Signature": signature}
                )
                durations.append(round(time.monotonic() - started, 3))
                response.raise_for_status()
                if cloud:
                    for play in ET.fromstring(response.text).iter("Play"):
                        started_audio = time.monotonic()
                        audio = await client.get(play.text)
                        audio.raise_for_status()
                        audio_seconds.append(round(time.monotonic() - started_audio, 3))
                        if audio.content[:4] != b"RIFF" or audio.content[8:12] != b"WAVE":
                            raise RuntimeError("Jessica audio was not valid WAV")
                        played_audio.append(audio.content)
                        with wave.open(io.BytesIO(audio.content), "rb") as wav:
                            playback_seconds = wav.getnframes() / wav.getframerate()
                        await asyncio.sleep(min(playback_seconds, 3))
                return response.text

            xml = await send(f"/twiml?session_id={session['id']}")
            answers = [
                "Yes, go ahead",
                "Finished my Python project",
                "Learned how webhooks work",
                "I spent 120 rupees on lunch and 80 rupees on transport",
                "No income today",
                "I am feeling good today",
                "Tomorrow I will finish the voice agent",
                "Nothing else",
                "Yes, that's correct",
            ]
            for speech in answers:
                gather = ET.fromstring(xml).find("Gather")
                if gather is None:
                    raise RuntimeError("Phone flow ended before all questions were verified")
                action = urlsplit(gather.attrib["action"])
                xml = await send(action.path + "?" + action.query, speech)
                if "<Stream" in xml:
                    raise RuntimeError("Streaming appeared in the trial phone flow")
        state = runtime.checkin.state(session["id"])
        if (
            state["status"] != "completed"
            or state["spending_totals"] != {"INR": "200.00"}
            or state["income_totals"]
        ):
            raise RuntimeError(
                f"Local check-in verification failed: status={state['status']}, spending={state['spending_totals']}, income={state['income_totals']}"
            )
        if max(durations) >= 5:
            raise RuntimeError("A webhook response exceeded Twilio's five-second trial deadline")
        runtime.store.update(session["id"], call_status="completed")
        result = await runtime.sheets.flush()
        if result.get("error"):
            raise RuntimeError("Verification workbook could not be saved")
        book = load_workbook(settings.excel_path, read_only=True)
        row = next(
            row
            for row in book[settings.google_sheet_tab].iter_rows(min_row=2, values_only=True)
            if row[0] == session["id"]
        )
        if row[2] != "completed":
            raise RuntimeError("Excel did not contain the completed check-in")
        book.close()
        receipt = {
            "verified": True,
            "model": settings.llm_model if cloud else settings.local_model,
            "voice": settings.tts_voice if cloud else settings.phone_voice,
            "voice_model": settings.tts_model if cloud else None,
            "checkin_id": session["id"],
            "spending_totals": state["spending_totals"],
            "webhook_seconds": durations,
            "max_webhook_seconds": max(durations),
            "excel": str(settings.excel_path),
            "phone_called": False,
            "audio_fetch_seconds": audio_seconds,
            "jessica_audio_responses": len(played_audio),
        }
        if cloud:
            (output / "Jessica Greeting.wav").write_bytes(played_audio[0])
        (output / "receipt.json").write_text(json.dumps(receipt, indent=2), encoding="utf-8")
        print(json.dumps(receipt, indent=2))
    finally:
        runtime.store.update(session["id"], call_status="completed")
        await runtime.close()


if __name__ == "__main__":
    asyncio.run(verify())
