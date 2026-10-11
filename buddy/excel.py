import asyncio
import os
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo
from buddy.storage import Store
from config import Settings

LOG_HEADERS = ["Log ID", "Check-in ID", "Timestamp", "Speaker", "Message"]
CHECKINS_TAB = "Buddy Checkins"
LOG_TAB = "Conversation Log"

class ExcelClient:
    label = "Excel file"
    configured = True

    def __init__(self, settings: Settings):
        self.settings = settings
        self.path = settings.excel_path.resolve()

    @property
    def target(self) -> str:
        return f"excel:{self.path}"

    @staticmethod
    def make_sheet(workbook, title: str, headers: list, widths: list):
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        sheet = workbook.create_sheet(title)
        sheet.append(headers)
        sheet.freeze_panes = "C2"
        sheet.sheet_view.showGridLines = False
        sheet.row_dimensions[1].height = 32
        for col, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(col)].width = width
        for cell in sheet[1]:
            cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="244A40")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        return sheet
