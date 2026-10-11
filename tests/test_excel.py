import asyncio
import json
import os
from threading import Event

import pytest
from openpyxl import load_workbook

from buddy.excel import CHECKINS_TAB, LOG_TAB, ExcelClient
from buddy.sheets import HEADERS, SheetsSync
from buddy.storage import Store


async def test_excel_records_answers_and_transcripts_without_google(settings):
    store = Store(settings.data_dir / "test.sqlite3")
    session = store.create("2026-10-12")
    store.update(
        session["id"],
        answers={"accomplishments": {"answer": "=2+2"}},
        expenses=[{"description": "Lunch", "amount": "80.50", "currency": "INR"}],
    )
    store.transcript(session["id"], "user", "=HYPERLINK(unsafe)")
    client = ExcelClient(settings)
    sync = SheetsSync(store, client)
    assert (await sync.flush())["synced"] == 1
    workbook = load_workbook(settings.excel_path)
    assert [cell.value for cell in workbook[CHECKINS_TAB][1]] == HEADERS
    assert workbook[CHECKINS_TAB]["D2"].value == "=2+2"
    assert workbook[CHECKINS_TAB]["D2"].data_type == "s"
    assert json.loads(workbook[CHECKINS_TAB]["G2"].value) == {"INR": "80.50"}
    assert workbook[LOG_TAB]["E2"].value == "=HYPERLINK(unsafe)"
    assert workbook[LOG_TAB]["E2"].data_type == "s"
    assert workbook[CHECKINS_TAB].freeze_panes == "C2"
    workbook.close()
    store.transcript(session["id"], "assistant", "Anything else today?")
    assert (await sync.flush())["synced"] == 1
    workbook = load_workbook(settings.excel_path)
    assert workbook[CHECKINS_TAB].max_row == 2
    assert workbook[LOG_TAB].max_row == 3
    workbook.close()


async def test_locked_excel_retains_previous_file_and_retries(settings, monkeypatch):
    store = Store(settings.data_dir / "test.sqlite3")
    session = store.create("2026-10-12")
    client = ExcelClient(settings)
    sync = SheetsSync(store, client)
    await sync.flush()
    original = settings.excel_path.read_bytes()
    store.update(session["id"], answers={"learning": {"answer": "Built Buddy"}})
    store.transcript(session["id"], "user", "Built Buddy")
    replace = os.replace

    def locked(*args):
        raise PermissionError("File in use")

    monkeypatch.setattr(os, "replace", locked)
    assert "locked" in (await sync.flush())["error"]
    assert settings.excel_path.read_bytes() == original
    assert len(store.pending(client.target, float("inf"))) == 1
    assert list(settings.data_dir.glob("*.xlsx")) == [settings.excel_path]
    monkeypatch.setattr(os, "replace", replace)
    store.retry_sync()
    assert (await sync.flush())["synced"] == 1
    workbook = load_workbook(settings.excel_path)
    assert workbook[CHECKINS_TAB]["E2"].value == "Built Buddy"
    assert workbook[LOG_TAB].max_row == 2
    workbook.close()


async def test_missing_excel_rebuilds_saved_history(settings):
    store = Store(settings.data_dir / "test.sqlite3")
    client = ExcelClient(settings)
    sync = SheetsSync(store, client)
    for day in ("2026-10-11", "2026-10-12"):
        session = store.create(day)
        store.transcript(session["id"], "user", day)
    await sync.flush()
    settings.excel_path.unlink()
    await client.initialize(store)
    workbook = load_workbook(settings.excel_path)
    assert workbook[CHECKINS_TAB].max_row == 3
    assert workbook[LOG_TAB].max_row == 3
    workbook.close()


async def test_unrelated_excel_headers_are_not_overwritten(settings):
    store = Store(settings.data_dir / "test.sqlite3")
    client = ExcelClient(settings)
    await client.initialize(store)
    workbook = load_workbook(settings.excel_path)
    workbook[CHECKINS_TAB]["A1"] = "Unrelated data"
    workbook.save(settings.excel_path)
    workbook.close()
    store.create("2026-10-12")
    result = await SheetsSync(store, client).flush()
    assert "headers" in result["error"]
    workbook = load_workbook(settings.excel_path)
    assert workbook[CHECKINS_TAB]["A1"].value == "Unrelated data"
    workbook.close()


async def test_shutdown_waits_for_atomic_excel_write_before_retry(settings, monkeypatch):
    store = Store(settings.data_dir / "test.sqlite3")
    store.create("2026-10-12")
    client = ExcelClient(settings)
    sync = SheetsSync(store, client)
    started, release = Event(), Event()
    persist = client.persist

    def slow_write(*args):
        started.set()
        if not release.wait(timeout=3):
            raise RuntimeError("Write was not released")
        persist(*args)

    monkeypatch.setattr(client, "persist", slow_write)
    task = asyncio.create_task(sync.flush())
    try:
        assert await asyncio.to_thread(started.wait, 2)
        task.cancel()
        await asyncio.sleep(0)
        assert not task.done()
    finally:
        release.set()
    with pytest.raises(asyncio.CancelledError):
        await task
    monkeypatch.setattr(client, "persist", persist)
    assert (await sync.flush())["synced"] == 1
    workbook = load_workbook(settings.excel_path)
    assert workbook[CHECKINS_TAB].max_row == 2
    workbook.close()
